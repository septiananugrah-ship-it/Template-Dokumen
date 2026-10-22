"""Helper python-docx untuk dokumen Word yang rapi dan ramah ATS
(teks asli, tanpa text box; tata letak kolom memakai tabel tanpa garis)."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Urutan anak elemen menurut skema OOXML; Word menolak file bila urutannya salah.
ORDER = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
            "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
            "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
            "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
            "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
            "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish",
            "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
            "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath"],
}


def insert_ordered(parent, child):
    """Sisipkan child (atau ganti yang setipe) di posisi yang benar."""
    kind = parent.tag.split("}")[1]
    order = ORDER[kind]
    name = child.tag.split("}")[1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    idx = order.index(name)
    for i, el in enumerate(parent):
        en = el.tag.split("}")[1]
        if en in order and order.index(en) > idx:
            el.addprevious(child)
            return child
    parent.append(child)
    return child


ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
         "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}


def rgb(h):
    return RGBColor.from_string(h)


def new_document(font="Calibri", size=10.5, margins=(1.8, 1.8, 1.8, 1.8), head_font=None,
                 color="222222", landscape=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin, sec.right_margin, sec.bottom_margin, sec.left_margin = (Cm(m) for m in margins)
    st = doc.styles["Normal"]
    st.font.name = font
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.font.color.rgb = rgb(color)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.1
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        hs = doc.styles[name]
        hs.font.name = head_font or font
        rpr = hs.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), head_font or font)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        hs.font.italic = False
        hs.paragraph_format.keep_with_next = True
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    lb = doc.styles["List Bullet"]
    lb.paragraph_format.space_after = Pt(1.5)
    lb.paragraph_format.left_indent = Cm(0.5)
    lb.paragraph_format.first_line_indent = Cm(-0.35)
    return doc


def content_width(doc):
    s = doc.sections[0]
    return (s.page_width - s.left_margin - s.right_margin) / 360000  # cm


def para(container, text="", size=None, bold=False, italic=False, color=None, align="l",
         before=0, after=0, font=None, style=None, line=None, caps=False, spacing=None,
         keep=False):
    """Tambah paragraf. text boleh str atau list (teks, opsi) untuk banyak run."""
    p = container.add_paragraph(style=style)
    p.alignment = ALIGN[align]
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    runs = [(text, {})] if isinstance(text, str) else text
    for t, opt in runs:
        add_run(p, t, size=opt.get("size", size), bold=opt.get("bold", bold),
                italic=opt.get("italic", italic), color=opt.get("color", color),
                font=opt.get("font", font), caps=opt.get("caps", caps),
                spacing=opt.get("spacing", spacing))
    return p


def add_run(p, text, size=None, bold=False, italic=False, color=None, font=None, caps=False,
            spacing=None):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = rgb(color)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if caps:
        r.font.all_caps = True
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        insert_ordered(r._element.get_or_add_rPr(), sp)
    return r


def heading(container, text, level=1, **kw):
    """Paragraf bergaya Heading (terbaca ATS & navigasi Word) dengan format kustom."""
    kw.setdefault("keep", True)
    return para(container, text, style=f"Heading {level}", **kw)


def border(p, side="bottom", color="999999", size=6, space=2):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(pPr, OxmlElement("w:pBdr"))
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)
    return p


def right_tab(p, pos_cm):
    p.paragraph_format.tab_stops.add_tab_stop(Cm(pos_cm), WD_TAB_ALIGNMENT.RIGHT)


def bullet(container, text, size=None, color=None, after=1.5, bold_lead=True):
    p = container.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(after)
    lead, sep, rest = text.partition(": ")
    if sep and bold_lead:
        add_run(p, lead + sep, size=size, bold=True, color=color)
        add_run(p, rest, size=size, color=color)
    else:
        add_run(p, text, size=size, color=color)
    return p


# ---------------------------------------------------------------- tabel tata letak

def layout_table(container, widths_cm, borders=False, cell_margin=(0.0, 0.0, 0.0, 0.0)):
    t = container.add_table(rows=1, cols=len(widths_cm))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    insert_ordered(tblPr, lay)
    if not borders:
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        insert_ordered(tblPr, b)
    mar = OxmlElement("w:tblCellMar")
    top, right, bottom, left = cell_margin
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(v * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    insert_ordered(tblPr, mar)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for c, w in zip(t.rows[0].cells, widths_cm):
        c.width = Cm(w)
    return t


def add_row(t, widths_cm):
    row = t.add_row()
    for c, w in zip(row.cells, widths_cm):
        c.width = Cm(w)
    return row


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd)


def cell_margins(cell, top=0, right=0, bottom=0, left=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(v * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    insert_ordered(tcPr, mar)


def valign(cell, where="top"):
    cell.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                               "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                               "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[where]


def cell_borders(cell, color="BBBBBB", size=6, sides=("top", "left", "bottom", "right"),
                 val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        if side in sides:
            e.set(qn("w:val"), val)
            e.set(qn("w:sz"), str(size))
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    insert_ordered(tcPr, b)


def row_height(row, cm, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(h)


def first_para(cell):
    """Sel baru selalu punya satu paragraf kosong; pakai itu agar tidak ada baris kosong."""
    return cell.paragraphs[0]


def fill_cell_para(cell, *args, **kw):
    """Seperti para(), tetapi paragraf pertama sel dipakai ulang."""
    p0 = cell.paragraphs[0]
    if p0.text == "" and not p0.runs and not getattr(cell, "_used", False):
        cell._used = True
        new = para(cell, *args, **kw)
        p0._p.getparent().remove(p0._p)
        return new
    return para(cell, *args, **kw)


def remove_leading_empty(cell):
    p0 = cell.paragraphs[0]
    if not p0.runs and len(cell.paragraphs) > 1:
        p0._p.getparent().remove(p0._p)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_core(doc, title, author="Template-Dokumen"):
    doc.core_properties.title = title
    doc.core_properties.author = author


# ---------------------------------------------------------------- field, seksi, nomor halaman

def add_field(p, instr, result="", size=None, bold=False, color=None):
    """Field Word (PAGE, TOC, ...). `result` = teks tampilan sebelum diperbarui."""
    def fld(kind):
        r = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            fc.set(qn("w:dirty"), "true")
        r.append(fc)
        return r
    p._p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    p._p.append(r)
    p._p.append(fld("separate"))
    if result:
        add_run(p, result, size=size, bold=bold, color=color)
    p._p.append(fld("end"))
    return p


def page_number_format(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        # pgNumType harus sebelum cols/docGrid dll.
        anchor = None
        for tag in ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                    "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid"):
            anchor = sectPr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(pg)
        else:
            sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))


def footer_page_number(section, size=11, align="c", color=None):
    section.footer.is_linked_to_previous = False
    ft = section.footer
    p = ft.paragraphs[0]
    p.alignment = ALIGN[align]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    add_field(p, "PAGE", "1", size=size, color=color)


def blank_footer(section):
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        for r in list(p.runs):
            r._element.getparent().remove(r._element)


def toc(container, entries, width_cm, size=12, line=1.5):
    """Daftar isi sebagai field TOC Word yang sudah terisi (entries = [(teks, level,
    halaman)]). Di Word: klik kanan > Update Field untuk memperbarui."""
    first = True
    for i, (txt, level, page) in enumerate(entries):
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = line
        pf.left_indent = Cm(0.8 * (level - 1))
        pf.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if first:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            r.append(fc)
            p._p.append(r)
            r = OxmlElement("w:r")
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-2" \\h \\z \\u '
            r.append(it)
            p._p.append(r)
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "separate")
            r.append(fc)
            p._p.append(r)
            first = False
        add_run(p, txt, size=size, bold=(level == 1))
        add_run(p, f"\t{page}", size=size)
        if i == len(entries) - 1:
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            r.append(fc)
            p._p.append(r)
