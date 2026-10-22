"""Helper dasar untuk menyusun slide dengan python-pptx.

Semua warna boleh ditulis sebagai nama slot tema ("accent1", "dk2", ...) atau
hex 6 digit ("1F1B4D"). Slot tema ditulis sebagai warna tema PowerPoint,
jadi pembeli bisa mengganti palet lewat Design > Variants > Colors.
"""
from __future__ import annotations

import copy
import io
import re
from functools import lru_cache
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ICON_DIR = ROOT / "assets" / "icons"

SLIDE_W = 13.333
SLIDE_H = 7.5

SLOTS = {
    "dk1": MSO_THEME_COLOR.DARK_1,
    "lt1": MSO_THEME_COLOR.LIGHT_1,
    "dk2": MSO_THEME_COLOR.DARK_2,
    "lt2": MSO_THEME_COLOR.LIGHT_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1,
    "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3,
    "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5,
    "accent6": MSO_THEME_COLOR.ACCENT_6,
}

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def _color_spec(color):
    """'accent1' | 'accent1+40' (lebih terang 40%) | 'accent1-25' | hex."""
    if color is None:
        return None
    m = re.fullmatch(r"([a-z]+\d?)([+-]\d+)?", color)
    if m and m.group(1) in SLOTS:
        return SLOTS[m.group(1)], (int(m.group(2)) / 100 if m.group(2) else 0.0)
    return RGBColor.from_string(color), 0.0


def set_color(color_format, color):
    value, brightness = _color_spec(color)
    if isinstance(value, RGBColor):
        color_format.rgb = value
    else:
        color_format.theme_color = value
        if brightness:
            color_format.brightness = brightness


def resolve_hex(theme, color):
    """Hex sebenarnya untuk sebuah spesifikasi warna (dipakai ikon PNG)."""
    m = re.fullmatch(r"([a-z]+\d?)([+-]\d+)?", color)
    if not (m and m.group(1) in SLOTS):
        return color
    base = theme.colors[m.group(1)]
    shift = int(m.group(2)) / 100 if m.group(2) else 0.0
    r, g, b = (int(base[i:i + 2], 16) for i in (0, 2, 4))
    if shift > 0:
        r, g, b = (round(c + (255 - c) * shift) for c in (r, g, b))
    elif shift < 0:
        r, g, b = (round(c * (1 + shift)) for c in (r, g, b))
    return f"{r:02X}{g:02X}{b:02X}"


# ---------------------------------------------------------------- presentasi

def new_presentation(theme):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    apply_theme(prs, theme)
    _tidy_layouts(prs, theme)
    return prs


def apply_theme(prs, theme):
    """Tulis palet & font tema ke ppt/theme/theme1.xml."""
    master = prs.slide_masters[0]
    part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    root.set("name", theme.name)
    scheme = root.find(f".//{{{a}}}clrScheme")
    scheme.set("name", theme.name)
    for slot, hexval in theme.colors.items():
        el = scheme.find(f"{{{a}}}{slot}")
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{a}}}srgbClr", val=hexval)
    fonts = root.find(f".//{{{a}}}fontScheme")
    fonts.set("name", theme.name)
    fonts.find(f"{{{a}}}majorFont/{{{a}}}latin").set("typeface", theme.head_font)
    fonts.find(f"{{{a}}}minorFont/{{{a}}}latin").set("typeface", theme.body_font)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _tidy_layouts(prs, theme):
    """Sisakan layout 'Blank' dan 'Title Only' agar menu New Slide rapi."""
    master = prs.slide_masters[0]
    bg = master.background.fill
    bg.solid()
    set_color(bg.fore_color, theme.bg)
    keep = {"Blank", "Title Only"}
    for layout in list(master.slide_layouts):
        if layout.name not in keep:
            master.slide_layouts.remove(layout)


def blank_slide(prs, theme, bg=None):
    layout = next(l for l in prs.slide_layouts if l.name == "Blank")
    slide = prs.slides.add_slide(layout)
    fill = slide.background.fill
    fill.solid()
    set_color(fill.fore_color, bg or theme.bg)
    return slide


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- bentuk

def _no_line(shape):
    shape.line.fill.background()


def shadow(shape, blur=18, dist=4, alpha=14, color="000000"):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(Pt(blur)), dist=str(Pt(dist)),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None,
          name=None, transparency=0, dash=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid()
        set_color(shp.fill.fore_color, fill)
        if transparency:
            _set_alpha(shp, 100 - transparency)
    else:
        shp.fill.background()
    if line:
        set_color(shp.line.color, line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = dash
    else:
        _no_line(shp)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    no_shadow(shp)
    style = shp._element.find(qn("p:style"))
    if style is not None:  # gaya bawaan memberi bayangan di sebagian aplikasi
        shp._element.remove(style)
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def _set_alpha(shp, opacity):
    clr = shp._element.spPr.find(qn("a:solidFill"))[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    etree.SubElement(clr, qn("a:alpha"), val=str(int(opacity * 1000)))


def rect(slide, x, y, w, h, fill=None, **kw):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, **kw)


def rrect(slide, x, y, w, h, fill=None, radius=0.15, **kw):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=radius, **kw)


def oval(slide, x, y, w, h, fill=None, **kw):
    return shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill, **kw)


def line(slide, x1, y1, x2, y2, color, width=1.0, dash=None, name=None, arrow=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    set_color(ln.line.color, color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    if arrow:
        lnel = ln.line._get_or_add_ln()
        etree.SubElement(lnel, qn("a:tailEnd"), type="triangle", w="med", len="med")
    if name:
        ln.name = name
    return ln


# ---------------------------------------------------------------- teks

def text(slide, x, y, w, h, content, size=16, color="dk1", bold=False, italic=False,
         font=None, align="l", anchor="t", spacing=None, line_spacing=1.1, para_after=0,
         name=None, margin=0, caps=False, shape_obj=None, autofit=False):
    """content: str (\n = paragraf baru) atau list paragraf; tiap paragraf str
    atau list run dict(text=..., bold=..., color=..., size=..., font=...)."""
    if shape_obj is None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        box = shape_obj
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if autofit else MSO_AUTO_SIZE.NONE
    m = Inches(margin) if not isinstance(margin, (tuple, list)) else None
    if m is not None:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    else:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Inches(v) for v in margin)
    tf.vertical_anchor = ANCHOR[anchor]
    paras = content.split("\n") if isinstance(content, str) else content
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if para_after:
            p.space_after = Pt(para_after)
        runs = [dict(text=para)] if isinstance(para, str) else para
        for r in runs:
            run = p.add_run()
            run.text = r["text"].upper() if (r.get("caps", caps)) else r["text"]
            f = run.font
            f.size = Pt(r.get("size", size))
            f.bold = r.get("bold", bold)
            f.italic = r.get("italic", italic)
            fname = r.get("font", font)
            if fname:
                f.name = fname
            set_color(f.color, r.get("color", color))
            sp = r.get("spacing", spacing)
            if sp:
                run._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
    if name:
        box.name = name
    return box


def bullets(slide, x, y, w, h, items, size=16, color="dk1", bullet_color="accent1",
            char="•", para_after=8, line_spacing=1.1, name=None, bold_lead=True):
    """Daftar berpoin dengan bullet asli PowerPoint (bukan karakter literal).
    Item boleh 'Judul: isi' -> bagian sebelum ':' ditebalkan."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(para_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.3)))
        pPr.set("indent", str(-Inches(0.3)))
        bc = etree.SubElement(pPr, qn("a:buClr"))
        val, br = _color_spec(bullet_color)
        if isinstance(val, RGBColor):
            etree.SubElement(bc, qn("a:srgbClr"), val=str(val))
        else:
            etree.SubElement(bc, qn("a:schemeClr"), val=_scheme_name(val))
        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(pPr, qn("a:buChar"), char=char)
        lead, sep, rest = item.partition(": ")
        parts = [(lead + sep, True), (rest, False)] if (sep and bold_lead) else [(item, False)]
        for txt, b in parts:
            run = p.add_run()
            run.text = txt
            run.font.size = Pt(size)
            run.font.bold = b
            set_color(run.font.color, color)
    if name:
        box.name = name
    return box


def _scheme_name(theme_color):
    return {v: k for k, v in SLOTS.items()}[theme_color]


# ---------------------------------------------------------------- ikon

@lru_cache(maxsize=512)
def icon_png(name, hexcolor, px=256, stroke=2.0):
    import cairosvg
    svg = (ICON_DIR / f"{name}.svg").read_text()
    svg = svg.replace('stroke="currentColor"', f'stroke="#{hexcolor}"')
    svg = svg.replace('stroke-width="2"', f'stroke-width="{stroke}"')
    return cairosvg.svg2png(bytestring=svg.encode(), output_width=px, output_height=px)


def icon(slide, theme, name, x, y, size, color="accent1", name_tag=None):
    png = icon_png(name, resolve_hex(theme, color))
    pic = slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(size), Inches(size))
    pic.name = name_tag or f"Ikon {name}"
    return pic


def picture(slide, data, x, y, w, h, name=None):
    pic = slide.shapes.add_picture(io.BytesIO(data), Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        pic.name = name
    return pic


# ---------------------------------------------------------------- grafik

def chart(slide, theme, kind, x, y, w, h, categories, series, colors=None, legend=None,
          labels=True, number_format='0', font_size=12, label_color=None, gap=80,
          axis_color=None, grid_color=None, smooth=False, label_pos=None, overlap=None,
          axis_min=None, axis_format=None):
    data = CategoryChartData(number_format=number_format)
    data.categories = categories
    for sname, values in series:
        data.add_series(sname, values)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.font.size = Pt(font_size)
    ch.font.name = theme.body_font
    set_color(ch.font.color, axis_color or theme.muted)
    ch.has_title = False
    if legend:
        ch.has_legend = True
        ch.legend.position = legend
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font_size)
    else:
        ch.has_legend = False
    colors = colors or ["accent1", "accent2", "accent3", "accent5", "accent6"]
    plot = ch.plots[0]
    is_pie = kind in (_pie_kinds())
    if hasattr(plot, "gap_width") and not is_pie:
        plot.gap_width = gap
        if overlap is not None:
            plot.overlap = overlap
    if is_pie:
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            set_color(point.format.fill.fore_color, colors[i % len(colors)])
            point.format.line.fill.solid()
            set_color(point.format.line.color, theme.bg)
            point.format.line.width = Pt(2)
    else:
        for i, s in enumerate(plot.series):
            c = colors[i % len(colors)]
            if "LINE" in kind.name:
                s.format.line.width = Pt(3)
                set_color(s.format.line.color, c)
                s.smooth = smooth
                from pptx.enum.chart import XL_MARKER_STYLE
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 8
                s.marker.format.fill.solid()
                set_color(s.marker.format.fill.fore_color, c)
                s.marker.format.line.fill.solid()
                set_color(s.marker.format.line.color, theme.bg)
            else:
                s.format.fill.solid()
                set_color(s.format.fill.fore_color, c)
                s.format.line.fill.background()
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_value = True
        dl.font.size = Pt(font_size)
        dl.font.bold = True
        set_color(dl.font.color, label_color or theme.text)
        dl.number_format = number_format
        dl.number_format_is_linked = False
        if label_pos is not None:
            dl.position = label_pos
    if not is_pie:
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.width = Pt(0.75)
        set_color(va.major_gridlines.format.line.color, grid_color or theme.grid)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(font_size)
        if axis_min is not None:
            va.minimum_scale = axis_min
        if axis_format:
            va.tick_labels.number_format = axis_format
            va.tick_labels.number_format_is_linked = False
        ca = ch.category_axis
        ca.format.line.width = Pt(1)
        set_color(ca.format.line.color, grid_color or theme.grid)
        ca.tick_labels.font.size = Pt(font_size)
        from pptx.enum.chart import XL_TICK_MARK
        ca.major_tick_mark = XL_TICK_MARK.NONE
        va.major_tick_mark = XL_TICK_MARK.NONE
    gf.name = "Grafik (klik kanan > Edit Data)"
    return gf


def _pie_kinds():
    from pptx.enum.chart import XL_CHART_TYPE as T
    return (T.PIE, T.DOUGHNUT, T.PIE_EXPLODED, T.DOUGHNUT_EXPLODED)


# ---------------------------------------------------------------- tabel

def table(slide, theme, x, y, w, rows, col_widths=None, row_h=0.5, header_fill="accent1",
          header_color="lt1", body_color=None, band_fill=None, body_fill=None, size=13,
          header_size=13, line_color=None, align=None, bold_first_col=False, font=None):
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gf.table
    # buang gaya bawaan agar warna kita yang tampil
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    col_widths = col_widths or [w / ncols] * ncols
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            head = r == 0
            fill = header_fill if head else (band_fill if (band_fill and r % 2 == 0) else body_fill)
            if fill:
                cell.fill.solid()
                set_color(cell.fill.fore_color, fill)
            else:
                cell.fill.background()
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN[(align[c] if align else "l")]
            run = p.add_run()
            run.text = str(rows[r][c])
            run.font.size = Pt(header_size if head else size)
            run.font.bold = head or (bold_first_col and c == 0)
            if font:
                run.font.name = font
            set_color(run.font.color, header_color if head else (body_color or theme.text))
            _cell_borders(cell, line_color or theme.grid)
    gf.name = "Tabel"
    return gf


def _cell_borders(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    val, _ = _color_spec(color)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag), w=str(Pt(0.75)), cap="flat", cmpd="sng", algn="ctr")
        if tag in ("a:lnL", "a:lnR"):
            etree.SubElement(ln, qn("a:noFill"))
            continue
        sf = etree.SubElement(ln, qn("a:solidFill"))
        if isinstance(val, RGBColor):
            etree.SubElement(sf, qn("a:srgbClr"), val=str(val))
        else:
            etree.SubElement(sf, qn("a:schemeClr"), val=_scheme_name(val))
    # urutan skema: garis harus sebelum isi (fill)
    fills = [el for el in tcPr if el.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f)
        tcPr.append(f)
