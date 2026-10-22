"""Tata letak slide isi, dipakai bersama oleh semua tema dan semua paket.

Builder menerima daftar spesifikasi slide (dict dengan kunci 'type') lalu
memanggil metode bernama sama. Tiap metode hanya memakai peran warna tema
(theme.accent, theme.surface, ...), sehingga satu konten tampil berbeda di
tiap tema.
"""
from __future__ import annotations

from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE

from . import core as k

X0, X1 = 0.7, 12.63          # batas kiri/kanan area konten
W = X1 - X0
Y_END = 6.6                  # batas bawah area konten


class Builder:
    def __init__(self, theme, meta):
        self.t = theme
        self.meta = meta
        self.prs = k.new_presentation(theme)
        self.page = 0

    # ------------------------------------------------------------ dasar
    def build(self, specs, path):
        for spec in specs:
            spec = dict(spec)
            kind = spec.pop("type")
            self.page += 1
            slide = getattr(self, kind)(**spec)
            if spec.get("notes"):
                k.notes(slide, spec["notes"])
        self.prs.core_properties.title = self.meta
        self.prs.core_properties.author = "Template-Dokumen"
        self.prs.save(path)
        return path

    def _content(self, title, kicker=None, notes=None):
        slide = k.blank_slide(self.prs, self.t)
        y = self.t.frame(slide, title, kicker, self.page, self.meta)
        return slide, y

    def _panel(self):
        """Tema dengan kartu warna-warni memakai panel netral di belakang grafik
        agar batang grafik tidak menyatu dengan warna kartu."""
        return "lt1" if getattr(self.t, "neutral_panel", False) else None

    # ------------------------------------------------------------ halaman khusus
    def cover(self, notes=None, **d):
        s = k.blank_slide(self.prs, self.t, self.t.cover_bg)
        self.t.cover(s, d)
        return s

    def divider(self, num, title, subtitle="", notes=None):
        s = k.blank_slide(self.prs, self.t, self.t.divider_bg)
        self.t.divider(s, num, title, subtitle)
        return s

    def closing(self, notes=None, **d):
        s = k.blank_slide(self.prs, self.t, self.t.closing_bg)
        self.t.closing(s, d)
        return s

    # ------------------------------------------------------------ isi
    def agenda(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        rows = (len(items) + 1) // 2
        col_w = (W - 0.5) / 2
        row_h = (Y_END - y + 0.3) / rows
        for i, (head, desc) in enumerate(items):
            col, row = divmod(i, rows)
            x = X0 + col * (col_w + 0.5)
            yy = y + row * row_h
            t.card(s, x, yy, col_w, row_h - 0.3, i=i, name=f"Agenda {i + 1}")
            d = 0.62
            t.badge(s, x + 0.3, yy + (row_h - 0.3 - d) / 2, d, f"{i + 1:02d}", i=i, size=15)
            ch = row_h - 0.3
            k.text(s, x + 1.2, yy + ch / 2 - 0.42, col_w - 1.45, 0.42, head, size=19, bold=True,
                   color=t.primary, font="+mj-lt", anchor="b")
            k.text(s, x + 1.2, yy + ch / 2 + 0.04, col_w - 1.45, 0.45, desc, size=14,
                   color=t.muted)
        return s

    def text_stats(self, title, paragraphs, stats, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        t.body(s, X0, y + 0.05, 6.6, Y_END - y, paragraphs, size=17, line_spacing=1.3,
               para_after=12, name="Paragraf")
        n = len(stats)
        gap = 0.3
        h = (Y_END - y - gap * (n - 1)) / n
        x = 7.85
        for i, (val, label) in enumerate(stats):
            yy = y + i * (h + gap)
            t.card(s, x, yy, X1 - x, h, i=i, name=f"Statistik {i + 1}")
            k.text(s, x + 0.35, yy, 2.1, h, val, size=36, bold=True, color=t.accent,
                   anchor="m", font="+mj-lt")
            k.text(s, x + 2.45, yy, X1 - x - 2.7, h, label, size=15, color=t.text, anchor="m")
        return s

    def numbered(self, title, items, kicker=None, notes=None, lead=None):
        s, y = self._content(title, kicker)
        t = self.t
        if lead:
            t.body(s, X0, y - 0.15, W, 0.5, lead, size=15, color=t.muted)
            y += 0.5
        n = len(items)
        gap = 0.25
        h = (Y_END - y - gap * (n - 1)) / n
        for i, (head, desc) in enumerate(items):
            yy = y + i * (h + gap)
            t.card(s, X0, yy, W, h, i=i, name=f"Butir {i + 1}")
            d = min(0.8, h - 0.35)
            t.badge(s, X0 + 0.35, yy + (h - d) / 2, d, str(i + 1), i=i, size=22)
            k.text(s, X0 + 1.5, yy + 0.05, 2.9, h - 0.1, head, size=18, bold=True,
                   color=t.primary, anchor="m", font="+mj-lt")
            k.text(s, X0 + 4.5, yy + 0.05, W - 4.85, h - 0.1, desc, size=16, color=t.text,
                   anchor="m", line_spacing=1.2)
        return s

    def two_col(self, title, cols, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        cw = (W - 0.4) / 2
        for i, (ic, head, items) in enumerate(cols):
            x = X0 + i * (cw + 0.4)
            t.card(s, x, y, cw, Y_END - y, i=i, name=f"Kolom {i + 1}")
            t.icon_disc(s, ic, x + 0.4, y + 0.4, 0.75, i=i)
            k.text(s, x + 1.35, y + 0.4, cw - 1.7, 0.75, head, size=22, bold=True,
                   color=t.primary, anchor="m", font="+mj-lt")
            k.bullets(s, x + 0.45, y + 1.5, cw - 0.9, Y_END - y - 1.7, items, size=16,
                      color=t.text, bullet_color=t.accent, para_after=14, line_spacing=1.2)
        return s

    def icon_grid(self, title, items, kicker=None, notes=None, lead=None):
        s, y = self._content(title, kicker)
        t = self.t
        if lead:
            t.body(s, X0, y - 0.15, W, 0.5, lead, size=15, color=t.muted)
            y += 0.55
        cols = 2 if len(items) == 4 else 3
        rows = (len(items) + cols - 1) // cols
        gap = 0.3
        cw = (W - gap * (cols - 1)) / cols
        ch = (Y_END - y - gap * (rows - 1)) / rows
        for i, (ic, head, desc) in enumerate(items):
            r, c = divmod(i, cols)
            x, yy = X0 + c * (cw + gap), y + r * (ch + gap)
            t.card(s, x, yy, cw, ch, i=i, name=f"Kartu {i + 1}")
            d = 0.8
            t.icon_disc(s, ic, x + 0.4, yy + 0.4, d, i=i)
            k.text(s, x + 1.5, yy + 0.4, cw - 1.8, 0.45, head, size=20, bold=True,
                   color=t.primary, font="+mj-lt")
            k.text(s, x + 1.5, yy + 0.9, cw - 1.8, ch - 1.0, desc, size=16, color=t.text,
                   line_spacing=1.2)
        return s

    def table_slide(self, title, rows, col_widths=None, kicker=None, notes=None, note=None,
                    align=None, row_h=None):
        s, y = self._content(title, kicker)
        t = self.t
        row_h = row_h or min(0.72, (Y_END - y - (0.6 if note else 0)) / len(rows))
        widths = col_widths or [W / len(rows[0])] * len(rows[0])
        scale = W / sum(widths)
        widths = [w * scale for w in widths]
        head_fill = t.accent if not t.dark else "dk2"
        head_color = t.on_accent if not t.dark else "accent1"
        band = None if t.dark else t.surface
        k.table(s, t, X0, y, W, rows, col_widths=widths, row_h=row_h, header_fill=head_fill,
                header_color=head_color, band_fill=band, size=14, header_size=14,
                line_color=t.grid, align=align, bold_first_col=True,
                body_fill=("dk1" if t.dark else None))
        if note:
            k.text(s, X0, Y_END - 0.35, W, 0.35, note, size=11, italic=True, color=t.muted)
        return s

    def flow(self, title, inputs, output, labels, caption, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(inputs)
        bw, bh, gap = 3.9, 1.25, 0.45
        total = n * bh + (n - 1) * gap
        top = y + (Y_END - 0.9 - y - total) / 2
        ox, ow, oh = 8.1, 4.1, 1.8
        oy = top + total / 2 - oh / 2
        for i, (head, sub) in enumerate(inputs):
            yy = top + i * (bh + gap)
            t.card(s, X0 + 0.4, yy, bw, bh, i=i, name=f"Variabel {i + 1}")
            k.text(s, X0 + 0.7, yy + 0.2, bw - 0.6, 0.45, head, size=18, bold=True,
                   color=t.primary, font="+mj-lt")
            k.text(s, X0 + 0.7, yy + 0.67, bw - 0.6, 0.4, sub, size=14, color=t.muted)
            sx, sy = X0 + 0.4 + bw, yy + bh / 2
            k.line(s, sx + 0.1, sy, ox - 0.12, oy + oh / 2, t.accent, 2.0, arrow=True,
                   name=f"Panah {labels[i]}")
            mx = (sx + ox) / 2
            my = (sy + oy + oh / 2) / 2
            t.badge(s, mx - 0.3, my - 0.3, 0.6, labels[i], i=i, size=13)
        k.rrect(s, ox, oy, ow, oh, t.accent if not t.dark else "dk2", radius=t.card_radius,
                name="Variabel terikat", line=("accent1" if t.dark else None))
        k.text(s, ox + 0.3, oy + 0.25, ow - 0.6, 0.6, output[0], size=20, bold=True,
               color=t.on_accent if not t.dark else "accent1", font="+mj-lt", align="c")
        k.text(s, ox + 0.3, oy + 0.85, ow - 0.6, 0.5, output[1], size=14,
               color=t.on_accent if not t.dark else "lt1", align="c")
        k.text(s, X0, Y_END - 0.7, W, 0.6, caption, size=15, italic=True, color=t.muted,
               align="c")
        return s

    def steps(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(items)
        cw = W / n
        d = 1.0
        cy = y + 0.2
        k.line(s, X0 + cw / 2, cy + d / 2, X1 - cw / 2, cy + d / 2, t.grid, 2.0,
               dash=MSO_LINE_DASH_STYLE.DASH, name="Garis alur")
        for i, (ic, head, desc) in enumerate(items):
            cx = X0 + i * cw + cw / 2
            fill = t.accent if not t.dark else "dk2"
            k.oval(s, cx - d / 2, cy, d, d, fill, name=f"Tahap {i + 1}",
                   line=("accent1" if t.dark else t.bg), line_w=(1.25 if t.dark else 4))
            k.icon(s, t, ic, cx - 0.26, cy + 0.24, 0.52, t.on_accent if not t.dark else "accent1")
            k.text(s, cx - cw / 2 + 0.1, cy + d + 0.25, cw - 0.2, 0.3, f"TAHAP {i + 1}",
                   size=11, bold=True, color=t.accent, align="c", spacing=1.5)
            k.text(s, cx - cw / 2 + 0.1, cy + d + 0.6, cw - 0.2, 0.75, head, size=17, bold=True,
                   color=t.primary, align="c", font="+mj-lt", line_spacing=1.0)
            t.card(s, cx - cw / 2 + 0.12, cy + d + 1.45, cw - 0.24, Y_END - (cy + d + 1.45), i=i,
                   name=f"Keterangan {i + 1}")
            k.text(s, cx - cw / 2 + 0.3, cy + d + 1.6, cw - 0.6, Y_END - (cy + d + 1.75), desc,
                   size=14, color=t.text, align="c", line_spacing=1.15)
        return s

    def cards3(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(items)
        gap = 0.35
        cw = (W - gap * (n - 1)) / n
        h = Y_END - y
        for i, (ic, head, value, desc) in enumerate(items):
            x = X0 + i * (cw + gap)
            t.card(s, x, y, cw, h, i=i, name=f"Kartu {i + 1}")
            t.icon_disc(s, ic, x + 0.4, y + 0.4, 0.75, i=i)
            k.text(s, x + 0.4, y + 1.35, cw - 0.8, 0.4, head, size=13, bold=True, color=t.muted,
                   caps=True, spacing=1.5)
            k.text(s, x + 0.4, y + 1.75, cw - 0.8, 0.8, value, size=28, bold=True,
                   color=t.primary, font="+mj-lt")
            k.text(s, x + 0.4, y + 2.6, cw - 0.8, h - 2.8, desc, size=16, color=t.text,
                   line_spacing=1.2)
        return s

    def chart_bar(self, title, categories, series, insights, kicker=None, notes=None,
                  number_format="0.0", chart_title=None):
        s, y = self._content(title, kicker)
        t = self.t
        cw = 7.7
        t.card(s, X0, y, cw, Y_END - y, i=0, name="Panel grafik", fill=self._panel())
        if chart_title:
            k.text(s, X0 + 0.35, y + 0.2, cw - 0.7, 0.4, chart_title, size=14, bold=True,
                   color=t.primary)
        legend = XL_LEGEND_POSITION.TOP if len(series) > 1 else None
        k.chart(s, t, XL_CHART_TYPE.COLUMN_CLUSTERED, X0 + 0.2, y + 0.6, cw - 0.4,
                Y_END - y - 0.75, categories, series, colors=list(t.series), legend=legend,
                number_format=number_format, label_pos=XL_LABEL_POSITION.OUTSIDE_END,
                gap=70, overlap=-10, axis_format="0")
        x = X0 + cw + 0.35
        n = len(insights)
        gap = 0.3
        h = (Y_END - y - gap * (n - 1)) / n
        for i, (val, label) in enumerate(insights):
            yy = y + i * (h + gap)
            t.card(s, x, yy, X1 - x, h, i=i + 1, name=f"Temuan {i + 1}")
            k.text(s, x + 0.35, yy + 0.2, X1 - x - 0.7, 0.75, val, size=34, bold=True,
                   color=t.accent, font="+mj-lt")
            k.text(s, x + 0.35, yy + 0.95, X1 - x - 0.7, h - 1.05, label, size=15, color=t.text,
                   line_spacing=1.15)
        return s

    def chart_mix(self, title, pie, line, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        lw = 5.0
        h = Y_END - y
        t.card(s, X0, y, lw, h, i=0, name="Panel diagram lingkar", fill=self._panel())
        k.text(s, X0 + 0.35, y + 0.2, lw - 0.7, 0.4, pie["title"], size=14, bold=True,
               color=t.primary)
        k.chart(s, t, XL_CHART_TYPE.DOUGHNUT, X0 + 0.2, y + 0.65, lw - 0.4, h - 0.85,
                pie["categories"], [("Data", pie["values"])], colors=list(t.series),
                legend=XL_LEGEND_POSITION.BOTTOM, number_format='0"%"',
                label_color=("dk1" if t.dark else "lt1"))
        x = X0 + lw + 0.35
        rw = X1 - x
        t.card(s, x, y, rw, h, i=1, name="Panel grafik garis", fill=self._panel())
        k.text(s, x + 0.35, y + 0.2, rw - 0.7, 0.4, line["title"], size=14, bold=True,
               color=t.primary)
        k.chart(s, t, XL_CHART_TYPE.LINE_MARKERS, x + 0.2, y + 0.65, rw - 0.4, h - 0.85,
                line["categories"], line["series"], colors=list(t.series),
                legend=XL_LEGEND_POSITION.TOP, number_format="0", labels=False,
                axis_min=line.get("min"))
        return s

    def photos(self, title, captions, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(captions)
        gap = 0.35
        cw = (W - gap * (n - 1)) / n
        ph = Y_END - y - 0.6
        for i, cap in enumerate(captions):
            x = X0 + i * (cw + gap)
            box = k.rrect(s, x, y, cw, ph, t.surface if not t.dark else "dk2",
                          radius=t.card_radius, line=t.grid, line_w=1.0,
                          dash=MSO_LINE_DASH_STYLE.DASH,
                          name=f"Foto {i + 1} (Format Shape > Fill > Picture)")
            k.icon(s, t, "image-plus", x + cw / 2 - 0.4, y + ph / 2 - 0.65, 0.8, t.muted)
            k.text(s, x + 0.2, y + ph / 2 + 0.3, cw - 0.4, 0.6,
                   "Klik kanan kotak > Format Shape > Fill > Picture", size=11, color=t.muted,
                   align="c")
            k.text(s, x, y + ph + 0.15, cw, 0.4, cap, size=13, bold=True, color=t.text, align="c")
        return s

    def discussion(self, title, quote, source, points, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        qw = 5.3
        h = Y_END - y
        k.rrect(s, X0, y, qw, h, t.accent if not t.dark else "dk2", radius=t.card_radius,
                name="Kutipan", line=("accent1" if t.dark else None))
        qc = t.on_accent if not t.dark else "accent1"
        k.icon(s, t, "quote", X0 + 0.45, y + 0.45, 0.7, qc)
        k.text(s, X0 + 0.45, y + 1.35, qw - 0.9, h - 2.3, quote, size=21, bold=True, color=qc,
               font="+mj-lt", line_spacing=1.2)
        k.text(s, X0 + 0.45, Y_END - 0.85, qw - 0.9, 0.5, source, size=13,
               color=(t.on_accent if not t.dark else "lt1"))
        x = X0 + qw + 0.5
        k.bullets(s, x, y + 0.1, X1 - x, h - 0.1, points, size=18, color=t.text,
                  bullet_color=t.accent, para_after=22, line_spacing=1.25)
        return s

    def conclusion(self, title, left, right, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        cw = (W - 0.4) / 2
        for ci, (ic, head, items) in enumerate((left, right)):
            x = X0 + ci * (cw + 0.4)
            t.card(s, x, y, cw, Y_END - y, i=ci, name=head)
            t.icon_disc(s, ic, x + 0.4, y + 0.35, 0.7, i=ci)
            k.text(s, x + 1.3, y + 0.35, cw - 1.6, 0.7, head, size=22, bold=True,
                   color=t.primary, anchor="m", font="+mj-lt")
            yy = y + 1.35
            step = (Y_END - yy - 0.2) / len(items)
            for j, item in enumerate(items):
                t.badge(s, x + 0.45, yy + j * step + 0.02, 0.42, str(j + 1), i=ci, size=12)
                k.text(s, x + 1.1, yy + j * step, cw - 1.45, step - 0.1, item, size=16,
                       color=t.text, line_spacing=1.15)
        return s

    def references(self, title, refs, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        box = s.shapes.add_textbox(*(k.Inches(v) for v in (X0, y, W, Y_END - y)))
        tf = box.text_frame
        tf.word_wrap = True
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", 0)
        for i, ref in enumerate(refs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = k.Pt(14)
            p.line_spacing = 1.2
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(k.Inches(0.5)))
            pPr.set("indent", str(-k.Inches(0.5)))
            author, _, rest = ref.partition(". ")
            for txt, b in ((author + ". ", True), (rest, False)):
                r = p.add_run()
                r.text = txt
                r.font.size = k.Pt(16)
                r.font.bold = b
                k.set_color(r.font.color, t.text)
        box.name = "Daftar pustaka (format APA)"
        return s

    def hypotheses(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(items)
        gap = 0.3
        h = (Y_END - y - gap * (n - 1)) / n
        for i, (code, desc) in enumerate(items):
            yy = y + i * (h + gap)
            t.card(s, X0, yy, W, h, i=i, name=f"Hipotesis {code}")
            d = min(0.9, h - 0.3)
            t.badge(s, X0 + 0.35, yy + (h - d) / 2, d, code, i=i, size=18)
            k.text(s, X0 + 1.6, yy, W - 1.95, h, desc, size=16, color=t.text, anchor="m",
                   line_spacing=1.2)
        return s

    def gantt(self, title, months, rows, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        label_w = 3.4
        cols = len(months)
        cw = (W - label_w) / cols
        n = len(rows)
        rh = (Y_END - y - 0.5) / n
        for j, m in enumerate(months):
            k.text(s, X0 + label_w + j * cw, y, cw, 0.4, m, size=12, bold=True, color=t.muted,
                   align="c", anchor="m")
        for i, (name, start, end) in enumerate(rows):
            yy = y + 0.5 + i * rh
            if i % 2 == 0:
                k.rect(s, X0, yy, W, rh, t.surface if not t.dark else "dk2", name="Baris")
            k.text(s, X0 + 0.15, yy, label_w - 0.2, rh, name, size=15, color=t.text, anchor="m")
            bar = k.rrect(s, X0 + label_w + start * cw + 0.06, yy + rh * 0.22,
                          (end - start + 1) * cw - 0.12, rh * 0.56,
                          t.series[i % len(t.series)], radius=rh * 0.28, name=f"Durasi {name}")
        return s

    # ------------------------------------------------------------ tambahan paket kuliah
    def team(self, title, members, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(members)
        gap = 0.3
        cw = (W - gap * (n - 1)) / n
        h = Y_END - y
        for i, (name, nim, role) in enumerate(members):
            x = X0 + i * (cw + gap)
            t.card(s, x, y, cw, h, i=i, name=f"Anggota {i + 1}")
            d = min(1.5, cw - 0.8)
            t.icon_disc(s, "user", x + (cw - d) / 2, y + 0.45, d, i=i)
            k.text(s, x + 0.15, y + d + 0.7, cw - 0.3, 0.8, name, size=17, bold=True,
                   color=t.primary, align="c", font="+mj-lt", line_spacing=1.0, anchor="b")
            k.text(s, x + 0.15, y + d + 1.55, cw - 0.3, 0.35, nim, size=13, color=t.muted,
                   align="c")
            k.text(s, x + 0.15, y + d + 2.0, cw - 0.3, 0.4, role, size=13, bold=True,
                   color=t.accent, align="c", caps=True, spacing=1)
        return s

    def quotes(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(items)
        gap = 0.35
        cw = (W - gap * (n - 1)) / n
        h = Y_END - y
        for i, (quote, author) in enumerate(items):
            x = X0 + i * (cw + gap)
            t.card(s, x, y, cw, h, i=i, name=f"Definisi {i + 1}")
            k.icon(s, t, "quote", x + 0.4, y + 0.4, 0.6, t.accent)
            k.text(s, x + 0.4, y + 1.2, cw - 0.8, h - 2.2, quote, size=16, italic=True,
                   color=t.text, line_spacing=1.3)
            k.text(s, x + 0.4, Y_END - 0.85, cw - 0.8, 0.5, author, size=15, bold=True,
                   color=t.primary, font="+mj-lt")
        return s

    def timeline(self, title, items, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        n = len(items)
        cw = W / n
        mid = y + (Y_END - y) / 2
        k.line(s, X0, mid, X1, mid, t.grid, 3.0, name="Garis waktu")
        for i, (year, head, desc) in enumerate(items):
            cx = X0 + i * cw + cw / 2
            d = 0.36
            k.oval(s, cx - d / 2, mid - d / 2, d, d, t.series[i % len(t.series)],
                   line=t.bg, line_w=3, name=f"Titik {year}")
            up = i % 2 == 0
            ty = mid - 2.15 if up else mid + 0.45
            ty = mid - 2.3 if up else mid + 0.45
            k.text(s, cx - cw / 2 + 0.05, ty, cw - 0.1, 0.6, year, size=30, bold=True,
                   color=t.primary, align="c", font="+mj-lt")
            k.text(s, cx - cw / 2 + 0.05, ty + 0.65, cw - 0.1, 0.4, head, size=17, bold=True,
                   color=t.accent if not getattr(t, "FILLS", None) else t.primary, align="c")
            k.text(s, cx - cw / 2 + 0.05, ty + 1.08, cw - 0.1, 0.75, desc, size=14,
                   color=t.muted, align="c", line_spacing=1.1)
        return s

    def compare(self, title, left, right, kicker=None, notes=None):
        s, y = self._content(title, kicker)
        t = self.t
        gap = 0.9
        cw = (W - gap) / 2
        h = Y_END - y
        for ci, (ic, head, items) in enumerate((left, right)):
            x = X0 + ci * (cw + gap)
            t.card(s, x, y, cw, h, i=ci, name=head)
            t.icon_disc(s, ic, x + 0.4, y + 0.35, 0.7, i=ci)
            k.text(s, x + 1.3, y + 0.35, cw - 1.6, 0.7, head, size=22, bold=True,
                   color=t.primary, anchor="m", font="+mj-lt")
            k.bullets(s, x + 0.45, y + 1.4, cw - 0.9, h - 1.6, items, size=16, color=t.text,
                      bullet_color=t.series[ci], para_after=12, line_spacing=1.2)
        d = 0.8
        t.badge(s, X0 + cw + (gap - d) / 2, y + h / 2 - d / 2, d, "VS", size=18)
        return s
